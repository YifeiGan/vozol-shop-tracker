#!/usr/bin/env python3
"""Import great_lakes.xlsx sheet Joey into Joey's Firestore profile."""

from __future__ import annotations

import json
import re
import urllib.error
import urllib.parse
import urllib.request
import uuid
from collections import OrderedDict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

XLSX_PATH = Path("/Users/ganyifei/Desktop/非常重要的东西/vozol-shop-tracker/great_lakes.xlsx")
SHEET_NAME = "Joey"
TOKEN_PATH = "/Users/ganyifei/.config/configstore/firebase-tools.json"
PROJECT = "vozol-orlando-management"
OWNER_UID = "YB9uqAmC56WUWgVnp9uhAMyWzVn1"
TEAM_ID = "great_lakes"
CHAIN_MIN_STORES = 5
MAX_TRAFFIC_NOTES = 2
DEFAULT_CITY = "Chicago"
BASE = f"https://firestore.googleapis.com/v1/projects/{PROJECT}/databases/(default)/documents"

OAUTH_CLIENT_ID = "563584335869-fgrhgmd47bqnekij5i8b5pr03ho849e6.apps.googleusercontent.com"
OAUTH_CLIENT_SECRET = "j9iVZfS8kkCEFUPaAeJV0sAi"

MONTH_NAMES = {
    "jan": 1, "january": 1,
    "feb": 2, "february": 2,
    "mar": 3, "march": 3,
    "apr": 4, "april": 4,
    "may": 5,
    "jun": 6, "june": 6,
    "jul": 7, "july": 7,
    "aug": 8, "august": 8,
    "sep": 9, "sept": 9, "september": 9,
    "oct": 10, "october": 10,
    "nov": 11, "november": 11,
    "dec": 12, "december": 12,
}


def blank(value):
    if value is None:
        return ""
    if isinstance(value, float):
        text = f"{value:.10f}".rstrip("0").rstrip(".")
    elif isinstance(value, int) and not isinstance(value, bool):
        text = str(value)
    elif hasattr(value, "strftime"):
        text = value.strftime("%Y-%m-%d")
    else:
        text = str(value)
    text = text.replace("\ufeff", "").replace("\xa0", " ").replace("，", ",")
    text = re.sub(r"[ \t]+", " ", text)
    text = text.strip()
    if text in ("-", "—", "–", "--", "null", "None", "none"):
        return ""
    return text


def load_token():
    with open(TOKEN_PATH) as fh:
        refresh = json.load(fh)["tokens"]["refresh_token"]
    body = urllib.parse.urlencode({
        "client_id": OAUTH_CLIENT_ID,
        "client_secret": OAUTH_CLIENT_SECRET,
        "refresh_token": refresh,
        "grant_type": "refresh_token",
    }).encode()
    req = urllib.request.Request("https://oauth2.googleapis.com/token", data=body, method="POST")
    with urllib.request.urlopen(req) as resp:
        return json.loads(resp.read())["access_token"]


def request(method, url, token, payload=None):
    data = None if payload is None else json.dumps(payload).encode()
    req = urllib.request.Request(
        url,
        data=data,
        method=method,
        headers={
            "Authorization": f"Bearer {token}",
            "Content-Type": "application/json",
        },
    )
    try:
        with urllib.request.urlopen(req) as resp:
            raw = resp.read()
            return json.loads(raw) if raw else {}
    except urllib.error.HTTPError as exc:
        body = exc.read().decode()
        raise RuntimeError(f"{method} {url} -> {exc.code}\n{body[:2000]}") from exc


def sval(v):
    return {"stringValue": v or ""}


def bval(v):
    return {"booleanValue": bool(v)}


def nval():
    return {"nullValue": None}


def ival(v):
    return {"integerValue": str(int(v))}


def aval(values):
    if not values:
        return {"arrayValue": {}}
    return {"arrayValue": {"values": values}}


def tsval(dt):
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return {"timestampValue": dt.astimezone(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")}


def visit_dt(date_str):
    return datetime.strptime(date_str, "%Y-%m-%d").replace(hour=16, minute=0, tzinfo=timezone.utc)


def parse_all_dates(text):
    dates = []
    if hasattr(text, "strftime") and not isinstance(text, str):
        return [text.strftime("%Y-%m-%d")]
    raw = blank(text)
    if not raw:
        return []
    for match in re.finditer(
        r"\b(Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|"
        r"Jul(?:y)?|Aug(?:ust)?|Sep(?:t(?:ember)?)?|Oct(?:ober)?|Nov(?:ember)?|"
        r"Dec(?:ember)?)\s+(\d{1,2})(?:[./,\s-]+(\d{2,4}))?\b",
        raw,
        re.I,
    ):
        month = MONTH_NAMES[match.group(1).lower()]
        day = int(match.group(2))
        year = 2026
        if match.group(3):
            y = int(match.group(3))
            year = 2000 + y if y < 100 else y
        if 1 <= day <= 31:
            value = f"{year:04d}-{month:02d}-{day:02d}"
            if value not in dates:
                dates.append(value)
    for match in re.finditer(r"(\d{4})[./-](\d{1,2})[./-](\d{1,2})", raw):
        year, month, day = int(match.group(1)), int(match.group(2)), int(match.group(3))
        if 1 <= month <= 12 and 1 <= day <= 31:
            value = f"{year:04d}-{month:02d}-{day:02d}"
            if value not in dates:
                dates.append(value)
    for match in re.finditer(r"(?<!\d)(\d{1,2})[./](\d{1,2})(?:[./](\d{2,4}))?", raw):
        month, day = int(match.group(1)), int(match.group(2))
        if not (1 <= month <= 12 and 1 <= day <= 31):
            continue
        year = 2026
        if match.group(3):
            y = int(match.group(3))
            year = 2000 + y if y < 100 else y
        value = f"{year:04d}-{month:02d}-{day:02d}"
        if value not in dates:
            dates.append(value)
    return dates


def parse_visit_date(text):
    dates = parse_all_dates(text)
    return dates[-1] if dates else ""


def fragment_date(text, fallback):
    dates = parse_all_dates(text)
    if dates and re.match(r"^[\*\-\s]*(\d{4}[./-])?\d{1,2}[./]\d{1,2}", blank(text)):
        return dates[0]
    if dates and re.match(r"^[\*\-\s]*(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)", blank(text), re.I):
        return dates[0]
    return fallback


def split_note_blocks(info, visit_date):
    text = blank(info)
    if not text:
        return []
    chunks = re.split(
        r"(?=\n\s*[\*\-]?\s*(?:\d{4}[./-])?\d{1,2}[./]\d{1,2})"
        r"|(?=\n\s*[\*\-]?\s*(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec))",
        text,
        flags=re.I,
    )
    blocks = []
    for chunk in chunks:
        chunk = chunk.strip()
        if not chunk:
            continue
        date = fragment_date(chunk, visit_date or "")
        blocks.append({"date": date, "text": chunk})
    if not blocks:
        return []
    if len(blocks) == 1 and not blocks[0]["date"]:
        blocks[0]["date"] = visit_date
    return blocks[-MAX_TRAFFIC_NOTES:]


def normalize_city(address):
    addr = blank(address)
    match = re.search(r",\s*([^,]+),\s*(IL|IN|MI|OH|WI|MN)\s*\d{0,5}", addr, re.I)
    if not match:
        match = re.search(r",\s*([^,]+),\s*(Illinois|Indiana|Michigan|Ohio|Wisconsin|Minnesota)\b", addr, re.I)
    if match:
        city = blank(match.group(1))
        if city and not re.search(r"\d", city) and 2 < len(city) < 40:
            return city
    return DEFAULT_CITY


def parse_store_count(value):
    match = re.search(r"(\d+)", blank(value))
    return int(match.group(1)) if match else 1


def parse_tier(text):
    text = blank(text).upper().replace("＋", "+")
    text = re.sub(r"\s+", "", text)
    if text in ("S", "A+", "A", "B"):
        return text
    if text.startswith("S+") or text.startswith("S＋"):
        return "S"
    return ""


def is_yes(text):
    value = blank(text).lower()
    return value in ("是", "yes", "y", "true", "1", "试抽盒", "样机")


def map_status(sell_status):
    value = blank(sell_status).replace("\n", "")
    if "已卖进" in value:
        return "visited"
    if "拒绝" in value:
        return "no_interest"
    return "follow_up"


def name_key(name):
    text = blank(name).lower().replace("&", " and ")
    text = re.sub(r"\band\b", "", text)
    return re.sub(r"[^a-z0-9]+", "", text)


def similar_name(a, b):
    na, nb = name_key(a), name_key(b)
    if not na or not nb:
        return False
    return na == nb or na in nb or nb in na


def addr_key(addr):
    return re.sub(r"[^a-z0-9]+", "", blank(addr).lower())


def join_notes(*parts):
    chunks = []
    for part in parts:
        text = blank(part)
        if text and text not in chunks:
            chunks.append(text)
    return "\n".join(chunks)


def coop_starred(coop):
    value = blank(coop)
    return "极高" in value or value.lower() in ("high", "极高", "very high")


def normalize_raw_row(raw):
    cols = list(raw) + [None] * max(0, 16 - len(raw))
    person = blank(cols[0])
    if person.lower() != "joey":
        return None
    name = blank(cols[1])
    address = blank(cols[2])
    if not name and not address:
        return None
    if not name and address:
        street = address.split(",")[0].strip()
        name = street or address
    return {
        "name": name,
        "address": address,
        "city": normalize_city(address),
        "shop_type": blank(cols[3]) or "smoke shop",
        "visit_raw": cols[4],
        "store_count": cols[5],
        "tier": cols[6],
        "owner_name": cols[7],
        "phone": cols[8],
        "sell_status": cols[9],
        "coop": cols[10],
        "notes": cols[11],
        "sample": cols[12],
        "test_case": cols[13],
        "distributor": cols[14],
        "popup": cols[15],
    }


def load_rows():
    wb = load_workbook(XLSX_PATH, data_only=True, read_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise SystemExit(f"Sheet {SHEET_NAME!r} not found: {wb.sheetnames}")
    ws = wb[SHEET_NAME]
    shops = []
    empty_streak = 0
    for raw in ws.iter_rows(min_row=2, values_only=True):
        if not raw or not any(raw):
            empty_streak += 1
            if empty_streak > 30:
                break
            continue
        empty_streak = 0
        parsed = normalize_raw_row(raw)
        if not parsed or not parsed.get("name"):
            continue
        visit_raw = parsed.get("visit_raw")
        visit_date = parse_visit_date(visit_raw)
        note_source = join_notes(parsed.get("notes"), parsed.get("popup"))
        shops.append({
            "name": parsed["name"],
            "address": parsed.get("address") or "",
            "city": parsed.get("city") or DEFAULT_CITY,
            "shop_type": parsed.get("shop_type") or "smoke shop",
            "visit_date": visit_date,
            "store_count": parse_store_count(parsed.get("store_count")),
            "tier": parse_tier(parsed.get("tier")),
            "owner_name": blank(parsed.get("owner_name")),
            "phone": blank(parsed.get("phone")),
            "sell_status": blank(parsed.get("sell_status")),
            "coop": blank(parsed.get("coop")),
            "info": note_source,
            "sample": is_yes(parsed.get("sample")),
            "test_case": is_yes(parsed.get("test_case")),
            "distributor": blank(parsed.get("distributor")),
            "notes": split_note_blocks(note_source, visit_date),
        })
    wb.close()
    return shops


def merge_shops(rows):
    merged = OrderedDict()
    for row in rows:
        addr = addr_key(row["address"])
        key = None
        if addr:
            for existing_key, existing in merged.items():
                if addr_key(existing["address"]) == addr and similar_name(existing["name"], row["name"]):
                    key = existing_key
                    break
        if key is None:
            key = str(uuid.uuid4())
            merged[key] = {**row, "id": key, "visits": []}
        shop = merged[key]
        later = (row["visit_date"] or "") >= (shop["visit_date"] or "")
        if later:
            for field in (
                "name", "address", "city", "shop_type", "visit_date", "store_count", "tier",
                "owner_name", "phone", "sell_status", "coop", "info", "distributor",
            ):
                if row.get(field):
                    shop[field] = row[field]
            shop["sample"] = shop["sample"] or row["sample"]
            shop["test_case"] = shop["test_case"] or row["test_case"]
        else:
            for field in ("city", "shop_type", "tier", "owner_name", "phone", "sell_status", "coop", "distributor"):
                if not shop.get(field) and row.get(field):
                    shop[field] = row[field]
            shop["sample"] = shop["sample"] or row["sample"]
            shop["test_case"] = shop["test_case"] or row["test_case"]
        shop["visits"].append(row)
    return merged


def note_maps(shop):
    blocks = []
    for visit in shop["visits"]:
        blocks.extend(visit.get("notes") or [])
    if not blocks and shop.get("info"):
        blocks = [{"date": shop.get("visit_date") or "", "text": shop["info"]}]
    seen = []
    unique = []
    for block in blocks:
        text = blank(block.get("text"))
        if not text or text in seen:
            continue
        seen.append(text)
        unique.append(block)
    unique = unique[-MAX_TRAFFIC_NOTES:]
    values = []
    for block in unique:
        date = block.get("date") or shop.get("visit_date") or ""
        at = visit_dt(date) if date else datetime.now(timezone.utc)
        values.append({
            "mapValue": {
                "fields": {
                    "date": sval(date),
                    "text": sval(block["text"]),
                    "at": ival(int(at.timestamp() * 1000)),
                }
            }
        })
    return values, unique


def shop_fields(shop):
    notes, unique = note_maps(shop)
    latest_date = shop.get("visit_date") or (unique[-1]["date"] if unique else "")
    created_date = shop["visits"][0].get("visit_date") or latest_date
    updated = visit_dt(latest_date) if latest_date else datetime.now(timezone.utc)
    created = visit_dt(created_date) if created_date else updated
    store_count = int(shop.get("store_count") or 1)
    is_chain = store_count >= CHAIN_MIN_STORES
    tier = shop.get("tier") or ""
    sample = bool(shop.get("sample"))
    test_case = bool(shop.get("test_case"))
    sample_on = latest_date if sample else ""
    test_on = latest_date if test_case else ""
    traffic_text = unique[-1]["text"] if unique else (shop.get("info") or "")
    status = map_status(shop.get("sell_status"))
    city = shop.get("city") or normalize_city(shop.get("address") or "")
    return {
        "name": sval(shop["name"]),
        "address": sval(shop.get("address") or ""),
        "city": sval(city),
        "phone": sval(shop.get("phone") or ""),
        "tier": sval(tier),
        "status": sval(status),
        "is_chain": bval(is_chain),
        "chain_name": sval(""),
        "chain_total_stores": ival(store_count) if is_chain else nval(),
        "chain_a_plus_count": ival(1 if is_chain and tier in ("S", "A+", "A") else 0) if is_chain else nval(),
        "staff_contact": sval(""),
        "owner_name": sval(shop.get("owner_name") or ""),
        "owner_schedule": sval(""),
        "contact_role": sval(""),
        "store_number": sval("" if is_chain or store_count <= 1 else f"{store_count}家"),
        "restock_status": sval(shop.get("sell_status") or ""),
        "distributor": sval(shop.get("distributor") or ""),
        "test_case_placed": bval(test_case),
        "test_case_placed_on": sval(test_on),
        "sample_placed": bval(sample),
        "sample_placed_on": sval(sample_on),
        "traffic_note": sval(traffic_text),
        "traffic_notes": aval(notes),
        "units_log": aval([]),
        "brands_note": sval(""),
        "next_plan": sval(""),
        "next_plan_date": sval(""),
        "next_plan_time": sval(""),
        "source_url": sval(""),
        "starred": bval(coop_starred(shop.get("coop"))),
        "team_id": sval(TEAM_ID),
        "assigned_to": sval(OWNER_UID),
        "created_at": tsval(created),
        "updated_at": tsval(updated),
        "lat": nval(),
        "lng": nval(),
        "geocode_query": sval(""),
        "geocode_failed": bval(False),
    }, unique, created


def visit_fields(block, created):
    date = block.get("date") or ""
    at = visit_dt(date) if date else created
    return {
        "date": sval(date),
        "units": ival(0),
        "note": sval(block.get("text") or ""),
        "decision_maker": sval(""),
        "restock_status": sval(""),
        "test_case_placed": bval(False),
        "next_plan": sval(""),
        "created_at": tsval(at),
    }


def doc_name(*parts):
    return f"projects/{PROJECT}/databases/(default)/documents/" + "/".join(parts)


def main():
    token = load_token()
    profile = request("GET", f"{BASE}/profiles/{OWNER_UID}", token)
    fields = profile.get("fields") or {}
    name = (fields.get("full_name") or {}).get("stringValue", "")
    team = (fields.get("team_id") or {}).get("stringValue", "")
    print(f"Owner: {name} ({OWNER_UID}) team={team}")

    existing = request("GET", f"{BASE}/profiles/{OWNER_UID}/shops?pageSize=300", token)
    existing_docs = existing.get("documents") or []
    print(f"Existing shops before import: {len(existing_docs)}")
    if existing_docs:
        raise SystemExit("Joey already has shops; aborting to avoid duplicates.")

    rows = load_rows()
    shops = merge_shops(rows)
    writes = []
    counts = {"follow_up": 0, "visited": 0, "no_interest": 0, "not_visited": 0}
    samples = 0
    cases = 0
    starred = 0
    chains = 0

    for shop in shops.values():
        fields, unique, created = shop_fields(shop)
        status = fields["status"]["stringValue"]
        counts[status] = counts.get(status, 0) + 1
        if fields["sample_placed"]["booleanValue"]:
            samples += 1
        if fields["test_case_placed"]["booleanValue"]:
            cases += 1
        if fields["starred"]["booleanValue"]:
            starred += 1
        if fields["is_chain"]["booleanValue"]:
            chains += 1
        writes.append({
            "update": {
                "name": doc_name("profiles", OWNER_UID, "shops", shop["id"]),
                "fields": fields,
            }
        })
        for index, block in enumerate(unique):
            vid = f"{block.get('date') or 'undated'}_{index}"
            writes.append({
                "update": {
                    "name": doc_name("profiles", OWNER_UID, "shops", shop["id"], "visits", vid),
                    "fields": visit_fields(block, created),
                }
            })

    print(f"Raw rows: {len(rows)}; merged shops: {len(shops)}")
    print(
        f"Will write {len(shops)} shops "
        f"({counts['follow_up']} follow-up, {counts['visited']} sold-in, "
        f"{counts['no_interest']} no-interest, {cases} test case, {samples} sample, "
        f"{chains} chain, {starred} starred)"
    )
    print(f"Commit writes: {len(writes)}")

    for start in range(0, len(writes), 400):
        chunk = writes[start:start + 400]
        request("POST", f"{BASE}:commit", token, {"writes": chunk})
        print(f"Committed {start + 1}-{start + len(chunk)}")

    after = request("GET", f"{BASE}/profiles/{OWNER_UID}/shops?pageSize=300", token)
    print(f"Shops after import: {len(after.get('documents') or [])}")


if __name__ == "__main__":
    main()
